"""
File handling utilities for extracting text from PDF, DOCX, and TXT/MD files.

Uses pathlib for modern path handling and a dispatch table for clean extensibility.
"""
from __future__ import annotations

from pathlib import Path

import fitz  # PyMuPDF
import docx


# ---------------------------------------------------------------------------
# Individual format extractors
# ---------------------------------------------------------------------------

def _extract_from_txt(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8")


def _extract_pages_from_txt(file_path: Path) -> list[dict]:
    text = _extract_from_txt(file_path)
    return [{"page_number": 1, "text": text, "token_count": len(text.split())}]


def _extract_from_pdf(file_path: Path) -> str:
    text = ""
    with fitz.open(str(file_path)) as pdf_doc:
        for page in pdf_doc:
            text += page.get_text() + "\n"
    return text


def _extract_pages_from_pdf(file_path: Path) -> list[dict]:
    pages: list[dict] = []
    with fitz.open(str(file_path)) as pdf_doc:
        for page_num, page in enumerate(pdf_doc, start=1):
            text = page.get_text()
            pages.append({
                "page_number": page_num,
                "text": text,
                "token_count": len(text.split()),
            })
    return pages


def _extract_from_docx(file_path: Path) -> str:
    doc = docx.Document(str(file_path))
    return "\n".join(para.text for para in doc.paragraphs)


def _extract_pages_from_docx(file_path: Path) -> list[dict]:
    doc = docx.Document(str(file_path))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    pages: list[dict] = []
    chunk_size = 20  # heuristic: 20 paragraphs per "page"
    for i in range(0, len(paragraphs), chunk_size):
        page_text = "\n".join(paragraphs[i : i + chunk_size])
        pages.append({
            "page_number": (i // chunk_size) + 1,
            "text": page_text,
            "token_count": len(page_text.split()),
        })
    return pages


# ---------------------------------------------------------------------------
# Dispatch tables — add new formats by inserting a single row
# ---------------------------------------------------------------------------

_TEXT_EXTRACTORS: dict[str, callable] = {
    ".txt": _extract_from_txt,
    ".md": _extract_from_txt,
    ".pdf": _extract_from_pdf,
    ".docx": _extract_from_docx,
}

_PAGE_EXTRACTORS: dict[str, callable] = {
    ".txt": _extract_pages_from_txt,
    ".md": _extract_pages_from_txt,
    ".pdf": _extract_pages_from_pdf,
    ".docx": _extract_pages_from_docx,
}

SUPPORTED_EXTENSIONS = frozenset(_TEXT_EXTRACTORS.keys())


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def extract_raw_text(file_path: str) -> str:
    """Extract full text from a supported document file."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    extractor = _TEXT_EXTRACTORS.get(ext)
    if not extractor:
        raise ValueError(f"Unsupported file type: {ext} (supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))})")
    return extractor(path)


def extract_pages_raw(file_path: str) -> list[dict]:
    """Extract page-level text from a supported document file."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    extractor = _PAGE_EXTRACTORS.get(ext)
    if not extractor:
        raise ValueError(f"Unsupported file type: {ext} (supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))})")
    return extractor(path)
